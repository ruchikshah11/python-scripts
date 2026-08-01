"""
SYNOPSIS
    Builds a formatted PDF or Excel report from a CSV file - the output
    format is inferred from the --output file extension (.pdf or .xlsx).

DESCRIPTION
    Copy this file as the starting point for a new script, then:
    - Update REQUIRED_MODULES if you depend on different packages
    - Adjust the styling in build_excel_report()/build_pdf_report() to taste

EXAMPLE
    python build_report.py --input data.csv --output report.xlsx
    python build_report.py --input data.csv --output report.pdf --title "Sales Report"

NOTES
    Created by  : Ruchik Shah
    Created on  : 2026-07-30
    Modified by :
    Modified on :
    Version     : 1.0.0
"""

import argparse
import csv
import importlib
import logging
import sys
import time
import uuid
from datetime import datetime, timedelta
from pathlib import Path

#region Module Dependency Check
REQUIRED_MODULES = ["openpyxl", "reportlab"]

for module_name in REQUIRED_MODULES:
    try:
        importlib.import_module(module_name)
    except ImportError:
        print(f"Required module '{module_name}' is not installed. Install it with: pip install {module_name}")
        sys.exit(1)

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
#endregion

#region Global Variables
SCRIPT_FOLDER = Path(__file__).resolve().parent
LOGS_DIRECTORY = SCRIPT_FOLDER / "Logs"
LOG_FILE_NAME = LOGS_DIRECTORY / f"build_report_{datetime.now():%Y%m%d}.log"
PURGE_LOG_DAYS = 7

CORRELATION_ID = str(uuid.uuid4())[:11]
HEADER_COLOR = "4472C4"
#endregion

#region Logging Setup
def check_log_directory():
    LOGS_DIRECTORY.mkdir(parents=True, exist_ok=True)


def delete_old_logs(logger):
    if PURGE_LOG_DAYS <= 0:
        return
    try:
        logger.info("Deleting log files older than %s days", PURGE_LOG_DAYS)
        cutoff = datetime.now() - timedelta(days=PURGE_LOG_DAYS)
        for log_file in LOGS_DIRECTORY.glob("*.log"):
            if datetime.fromtimestamp(log_file.stat().st_mtime) < cutoff:
                log_file.unlink()
        logger.info("Log files deleted")
    except Exception as ex:
        logger.error("Error deleting log files. Details: %s", ex)


def get_logger():
    logger = logging.getLogger("build_report")
    logger.setLevel(logging.DEBUG)

    formatter = logging.Formatter("%(asctime)s %(levelname)s\t%(message)s", datefmt="%Y-%m-%d %H:%M:%S")

    file_handler = logging.FileHandler(LOG_FILE_NAME, encoding="utf-8")
    file_handler.setFormatter(formatter)

    console_handler = logging.StreamHandler()
    console_handler.setFormatter(logging.Formatter("%(message)s"))

    logger.addHandler(file_handler)
    logger.addHandler(console_handler)
    return logger
#endregion


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Path to the input CSV file")
    parser.add_argument("--output", required=True, help="Path to the output file - .xlsx or .pdf")
    parser.add_argument("--title", default="Report", help="Title shown at the top of the report")
    return parser.parse_args()


def read_csv(input_path):
    with open(input_path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))

    if not rows:
        raise ValueError(f"'{input_path}' is empty")

    return rows[0], rows[1:]


def build_excel_report(header, data_rows, output_path, title, logger):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Report"

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
    title_cell = worksheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center")

    header_row = 3
    for col_index, column_name in enumerate(header, start=1):
        cell = worksheet.cell(row=header_row, column=col_index, value=column_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=HEADER_COLOR, end_color=HEADER_COLOR, fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for row_offset, row_data in enumerate(data_rows, start=1):
        for col_index, value in enumerate(row_data, start=1):
            worksheet.cell(row=header_row + row_offset, column=col_index, value=value)

    for col_index, column_name in enumerate(header, start=1):
        max_length = len(str(column_name))
        for row_data in data_rows:
            if col_index - 1 < len(row_data):
                max_length = max(max_length, len(str(row_data[col_index - 1])))
        worksheet.column_dimensions[get_column_letter(col_index)].width = max_length + 4

    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)

    workbook.save(output_path)
    logger.info("Wrote %s data row(s) to Excel report: %s", len(data_rows), output_path)


def build_pdf_report(header, data_rows, output_path, title, logger):
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(str(output_path), pagesize=landscape(letter))

    table_data = [header] + [[str(value) for value in row] for row in data_rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{HEADER_COLOR}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))

    elements = [Paragraph(title, styles["Title"]), Spacer(1, 0.25 * inch), table]
    document.build(elements)
    logger.info("Wrote %s data row(s) to PDF report: %s", len(data_rows), output_path)


def main_process(args, logger):
    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.is_file():
        raise ValueError(f"Input file not found: {input_path}")

    header, data_rows = read_csv(input_path)
    logger.info("Read %s data row(s), %s column(s) from %s", len(data_rows), len(header), input_path)

    extension = output_path.suffix.lower()
    if extension == ".xlsx":
        build_excel_report(header, data_rows, output_path, args.title, logger)
    elif extension == ".pdf":
        build_pdf_report(header, data_rows, output_path, args.title, logger)
    else:
        raise ValueError(f"Unsupported output extension '{extension}' - use .xlsx or .pdf")


def main():
    args = parse_args()
    check_log_directory()
    logger = get_logger()

    script_start_time = datetime.now()
    stopwatch_start = time.perf_counter()

    try:
        delete_old_logs(logger)
        logger.info("=" * 64)
        logger.info("Script started at %s (CorrelationID: %s)", script_start_time, CORRELATION_ID)

        main_process(args, logger)

        elapsed = time.perf_counter() - stopwatch_start
        logger.info("Script finished at %s, time taken: %.3fs (CorrelationID: %s)",
                     datetime.now(), elapsed, CORRELATION_ID)
        logger.info("=" * 64)
    except Exception as ex:
        elapsed = time.perf_counter() - stopwatch_start
        logger.error("Script failed at %s, time taken: %.3fs (CorrelationID: %s). Details: %s",
                     datetime.now(), elapsed, CORRELATION_ID, ex)
        sys.exit(1)


if __name__ == "__main__":
    main()
